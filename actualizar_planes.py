"""Modelo completo del Power BI: fuentes, transformaciones, relaciones y cálculos.

Instalar una vez: python -m pip install pyodbc==5.3.0 polars==1.35.2 openpyxl==3.1.5
Ejecutar: python actualizar_planes.py
Requiere el DSN ODBC de 64 bits «salud» existente en la PC con acceso a IBM i.
Copiar también Despliegue 2022.xlsx, origen de tres tablas del Power BI.
Pide usuario y contraseña; solo ejecuta SELECT y escribe Parquet locales.

Desde Python: modelo = actualizar_modelo(excel='Despliegue 2022.xlsx')
Leer la última carga sin conectar: modelo = ModeloPowerBI('modelo_powerbi')
Consultar: modelo.medidas({'Planes Operativos': {'ANO': [2026]}})
Detalle: modelo.tabla('Estrategias').head(10).collect()
No necesita el PBIX para ejecutarse. La medida sin fórmula «prueba» queda bloqueada.
"""
from contextlib import closing, contextmanager
from datetime import date, datetime, time, timedelta, timezone
from getpass import getpass
import json
import hashlib
import inspect
import argparse
import os
from pathlib import Path
import posixpath
import re
import shutil
import tempfile
import sys
from time import perf_counter
import uuid
import zipfile
from xml.etree import ElementTree as ET

os.environ.setdefault('POLARS_MAX_THREADS', '2')
import polars as pl
import pyodbc

CONSULTA = 'SELECT * FROM JFM.VTPLANOPERATIVO'
CONSULTA_INCREMENTAL = CONSULTA + ' WHERE (ANO >= ? OR ANO IS NULL)'
# Colombia no cambia de hora; UTC-5 evita depender de tzdata en Windows.
COLOMBIA = timezone(timedelta(hours=-5), name='America/Bogota')
ESQUEMA = {
    'ID': pl.Float64, 'AREA': pl.String, 'PROCESO': pl.String, 'PROBOPR': pl.String,
    'ESTRATEG': pl.String, 'ACTIVID': pl.String, 'META': pl.String, 'TAREAS': pl.String,
    'FECINICIAL': pl.Datetime('ns'), 'FECFINAL': pl.Datetime('ns'),
    'RESPONSABLE': pl.String, 'RECURSOS': pl.String, 'INDICADOR': pl.String,
    'RESREAL': pl.String, 'EJECUCION': pl.Float64, 'CRONOGRAMA': pl.Float64,
    'LOGMETA': pl.Float64, 'FACTLIM': pl.String, 'IDENC': pl.String, 'PERIODO': pl.Int64,
    'ANO': pl.Int64, 'CALIFICACION': pl.Float64, 'ESTADO': pl.String, 'OBSESTADO': pl.String,
    'Trimestre': pl.String, 'Trimestre Incial': pl.Int64, 'Filtro': pl.String,
    'Año': pl.Int64, 'llave': pl.String, 'Proceso Final': pl.String,
}
PLAN, ESTRATEGIAS, MATRIZ, CORPORATIVAS = 'Planes Operativos', 'Estrategias', 'Matriz Objetivos', 'Estrategias Corporativas'
ESQUEMAS = {
    PLAN: ESQUEMA,
    ESTRATEGIAS: {
        **dict(list(ESQUEMA.items())[:28]), 'ID': pl.Int64,
        'EJECUCION': pl.String, 'CRONOGRAMA': pl.String, 'LOGMETA': pl.String, 'CALIFICACION': pl.String,
        'CRUCE ESTRATEGIAS': pl.String, 'Peso': pl.Float64, 'LLAVE': pl.Int64,
        'Calificación': pl.Float64, 'valor': pl.Float64, 'prueba': pl.Float64, 'LLAVE2': pl.String,
    },
    MATRIZ: {'Cod_Política': pl.Int64, 'Política': pl.String, 'Cod_Obj': pl.Int64,
        'Objetivo': pl.String, 'Area': pl.String, 'Cod_Estrategia': pl.Int64,
        'Estratégia': pl.String, 'Cuenta': pl.Int64, 'Peso': pl.Float64,
        'valor 1 trim': pl.Float64, 'Peso 1 trim': pl.Float64, 'VALOR': pl.Float64,
        'PESO2': pl.Float64, 'Trimestre': pl.String, 'Lllave': pl.String},
    CORPORATIVAS: {'Area': pl.String, 'Cod_Estrategia': pl.Int64, 'Estratégia': pl.String, 'Cruce': pl.String},
}
TRIMESTRES = ('1er Trimestre', '2do Trimestre', '3er Trimestre', '4to Trimestre')
METAS = dict(zip(TRIMESTRES, (.25, .5, .75, 1.)))
# Padre (uno) -> hijo (muchos); no se propagan filtros en sentido inverso.
RELACIONES = (
    (PLAN, 'IDENC', ESTRATEGIAS, 'IDENC'),
    (CORPORATIVAS, 'Cod_Estrategia', MATRIZ, 'Cod_Estrategia'),
    ('Trimestres', 'Trimestre', ESTRATEGIAS, 'Trimestre'),
    ('Trimestres', 'Trimestre', MATRIZ, 'Trimestre'),
    ('Metas', 'Trimestre', PLAN, 'Trimestre'),
)
CALENDARIOS = {
    'LocalDateTable_d0877634-0e73-4823-a111-776ba6bd19f7': (PLAN, 'FECINICIAL'),
    'LocalDateTable_e2b2ac52-1e62-41ee-89cd-3f4366ff3a0c': (PLAN, 'FECFINAL'),
    'LocalDateTable_0baccae0-ccd1-4113-b182-0cf7ef527e89': (ESTRATEGIAS, 'FECINICIAL'),
    'LocalDateTable_c9e851ec-cdba-4381-9d60-9223d73cbb04': (ESTRATEGIAS, 'FECFINAL'),
    'LocalDateTable_78b5bc63-b819-4ef1-b387-0ad792d53779': ('Fecha Actualización', 'Fecha Actualización'),
}
AREA_REPLACEMENTS = {'Apoyo Direccion': 'PROCESOS DE APOYO DIRECCIÓN', 'Administracion': 'ADMINISTRATIVA'}
PROCESS_REPLACEMENTS = {
    'INFRAESTRUCTURA': 'Infraestructura', 'MERCADEO ESTRATÉGICO': 'Mercadeo Estratégico',
    'CRÉDITO': 'Crédito', 'ÁREA DE CULTURA': 'Área de Cultura', 'IMPUESTOS': 'Impuestos',
    'CRECE EMPRESARIAL': 'Crece Empresarial', 'CISE SEDE LA VIRGINIA': 'CISE la Virginia',
    'CONTABILIDAD': 'Contabilidad', 'GESTIÓN FINANCIERA': 'Gestión Financiera',
}
FINAL_PROCESS = {
    'Agencia de Empelo / Fosfec': 'Agencia de Empleo / Fosfec',
    'Atecnion Integral A La Niñez': 'Atención Integral a la Niñez',
    'Desarrollo Fisico': 'Infraestructura', 'Administracion Jde': 'Sistemas Informáticos Financieros',
}


class ErrorDatos(ValueError):
    """Error de configuración o de datos que puede mostrarse sin credenciales."""


def _avance(mensaje):
    print(f'[{datetime.now():%H:%M:%S}] {mensaje}', flush=True)


def transform_operational(frame):
    """Reglas M y seis columnas DAX recuperadas; ANO y Año son campos distintos."""
    required = set(list(ESQUEMA)[:24])
    if required - set(frame.columns):
        raise ErrorDatos('La consulta no devolvió todas las columnas esperadas del Power BI.')
    date_col = pl.col('FECINICIAL')
    if frame.schema['FECINICIAL'] == pl.String:
        date_col = date_col.str.to_datetime(strict=True)
    frame = frame.with_columns(date_col.cast(pl.Datetime('ns')), pl.col('PERIODO').cast(pl.Int64),
                               pl.col('ID').cast(pl.Float64))
    frame = frame.filter((pl.col('ID').is_null() | ~pl.col('ID').is_in([7875, 7876])) &
                         (pl.col('ANO').is_null() | (pl.col('ANO') != 2027)) & pl.col('FECINICIAL').is_not_null())
    for column, replacements in [('AREA', AREA_REPLACEMENTS), ('PROCESO', PROCESS_REPLACEMENTS)]:
        expr = pl.col(column).cast(pl.String)
        for old, new in replacements.items():
            expr = expr.str.replace_all(old, new, literal=True)
        frame = frame.with_columns(expr.alias(column))
    frame = frame.with_columns(pl.when(pl.col('PROCESO') == 'Gerontología')
                               .then(pl.lit('DESARROLLO ESTRATÉGICO')).otherwise(pl.col('AREA')).alias('AREA'))
    quarter_no = pl.col('PERIODO').cast(pl.String).str.slice(-2).replace_strict(
        {'03': 1, '06': 2, '09': 3, '12': 4}, default=0, return_dtype=pl.Int64)
    quarters = {1: '1er Trimestre', 2: '2do Trimestre', 3: '3er Trimestre', 4: '4to Trimestre'}
    frame = frame.with_columns(
        quarter_no.alias('_quarter'),
        quarter_no.replace_strict(quarters, default='', return_dtype=pl.String).alias('Trimestre'),
        pl.col('FECINICIAL').dt.quarter().cast(pl.Int64).alias('Trimestre Incial'),
        pl.col('FECINICIAL').dt.year().cast(pl.Int64).alias('Año'),
        pl.col('PROCESO').str.to_lowercase().replace_strict(
            {key.lower(): value for key, value in FINAL_PROCESS.items()}, default=pl.col('PROCESO')).alias('Proceso Final'),
    )
    id_text = pl.when(pl.col('ID') % 1 == 0).then(pl.col('ID').cast(pl.Int64).cast(pl.String)).otherwise(pl.col('ID').cast(pl.String)).fill_null('')
    return frame.with_columns(
        pl.when((pl.col('_quarter') >= pl.col('Trimestre Incial')) & (pl.col('_quarter') > 0))
        .then(pl.lit('Aplica')).otherwise(pl.lit('No Aplica')).alias('Filtro'),
        (id_text + pl.col('PERIODO').cast(pl.String).fill_null('')).alias('llave'),
    ).drop('_quarter')


def convertir_modelo(frame, tabla=PLAN, desde_excel=False):
    """Conservar nombres, tipos, nulos y eliminación de espacios finales de Tabular."""
    schema = ESQUEMAS[tabla]
    if set(schema) - set(frame.columns):
        raise ErrorDatos(f'Faltan columnas esperadas en {tabla}: {sorted(set(schema) - set(frame.columns))}')
    if desde_excel:
        numeric = []
        for name, target in schema.items():
            if target not in {pl.Int64, pl.Float64}:
                continue
            column = pl.col(name)
            source = frame.schema[name]
            if source == pl.String:
                # Cultura es-CO recuperada: «5.1.1» se interpreta como 511.
                column = column.str.strip_chars().str.replace_all('.', '', literal=True).str.replace_all(',', '.', literal=True)
                if target == pl.Int64:
                    column = column.cast(pl.Decimal(38, 10), strict=True).round(0, mode='half_to_even')
                else:
                    column = column.cast(pl.Float64, strict=True)
            elif target == pl.Int64 and source in {pl.Float32, pl.Float64}:
                column = column.round(0, mode='half_to_even')
            numeric.append(column.cast(target, strict=True).alias(name))
        frame = frame.with_columns(numeric)
    expressions = []
    for name, dtype in schema.items():
        column = pl.col(name)
        if dtype == pl.Datetime('ns') and frame.schema[name] == pl.String:
            column = column.str.to_datetime(strict=True)
        column = column.cast(dtype, strict=True)
        if dtype == pl.String:
            column = column.str.strip_chars_end(' ')
        expressions.append(column.alias(name))
    return frame.select(expressions)


def extraer_y_guardar(cursor, salida, lote=10000, *, ano_desde=None, base=None, metricas=None):
    """Leer por lotes, ordenar por ANO en disco y publicar solo cuando termine."""
    if type(lote) is not int or lote < 1:
        raise ErrorDatos('El tamaño del lote debe ser un entero positivo.')
    target = Path(salida).expanduser().resolve()
    if target.suffix.lower() != '.parquet':
        raise ErrorDatos('El archivo de salida debe terminar en .parquet.')
    target.parent.mkdir(parents=True, exist_ok=True)
    metrics = metricas if metricas is not None else {}
    metrics.update(filas_descargadas=0, filas_ventana_publicadas=0,
                   filas_historicas_conservadas=0, consulta_s=0., descarga_s=0.)
    if (ano_desde is None) != (base is None):
        raise ErrorDatos('La actualización incremental requiere año y base completos.')
    total = final = 0
    with tempfile.TemporaryDirectory(prefix='.actualizando-', dir=target.parent) as folder:
        work = Path(folder).resolve()
        if not work.is_relative_to(target.parent):
            raise ErrorDatos('La carpeta temporal está fuera de la carpeta de salida.')
        query = CONSULTA if ano_desde is None else CONSULTA_INCREMENTAL
        metrics.update(sql=query, parametros=[] if ano_desde is None else [ano_desde])
        _avance(f'[3/3] Consulta de solo lectura: {query}; parámetros: {metrics["parametros"]}')
        started = perf_counter()
        if ano_desde is None:
            cursor.execute(query)
        else:
            cursor.execute(query, ano_desde)
        metrics['consulta_s'] = perf_counter() - started
        _avance('La consulta respondió. Esperando las primeras filas...')
        columns = [item[0] for item in cursor.description]
        if set(list(ESQUEMA)[:24]) - set(columns):
            raise ErrorDatos('La consulta no devolvió todas las columnas esperadas del Power BI.')
        parts = []
        while True:
            # Mostrar pronto el primer avance; el resto conserva el lote configurado.
            size = min(lote, 1000) if total == 0 else lote
            started = perf_counter()
            rows = cursor.fetchmany(size)
            metrics['descarga_s'] += perf_counter() - started
            if not rows:
                break
            total += len(rows)
            _avance(f'Filas recibidas: {total:,}. Procesando el lote...')
            frame = pl.DataFrame([tuple(row) for row in rows], schema=columns,
                                 orient='row', infer_schema_length=None)
            frame = convertir_modelo(transform_operational(frame))
            final += frame.height
            # Agrupar cada lote por año evita ordenar el conjunto completo en RAM.
            for (year,), group in frame.partition_by('ANO', as_dict=True, maintain_order=True).items():
                part = work / f'{len(parts):08d}.parquet'
                group.write_parquet(part, compression='zstd')
                parts.append((year, part))
        metrics.update(filas_descargadas=total, filas_ventana_publicadas=final)
        if not final and ano_desde is None:
            raise ErrorDatos('La consulta quedó sin filas después de los filtros. Se conserva la carga anterior.')
        _avance('Guardando y ordenando los datos locales...')
        parts.sort(key=lambda item: (item[0] is None, -(item[0] or 0)))
        ready = work / 'completo.parquet'
        # Reemplazo íntegro, incluidos NULL y una ventana sin registros.
        window = (pl.scan_parquet([path for _, path in parts], parallel='none')
                  if parts else pl.DataFrame(schema=ESQUEMA).lazy())
        if base is not None:
            historical = pl.scan_parquet(base).filter(pl.col('ANO') < ano_desde)
            metrics['filas_historicas_conservadas'] = historical.select(pl.len()).collect().item()
            window = pl.concat([window, historical])
        metrics['filas_totales_publicadas'] = final + metrics['filas_historicas_conservadas']
        window.sink_parquet(
            ready, compression='zstd', maintain_order=True, engine='streaming')
        os.replace(ready, target)
    _avance(f'Descargadas: {total:,}; ventana publicada: {final:,}; histórico conservado: {metrics["filas_historicas_conservadas"]:,}.')
    _avance(f'Archivo: {target}')
    return target


def actualizar_planes(dsn='salud', usuario=None, contrasena=None,
                      salida='planes_operativos.parquet', lote=10000, *,
                      ano_desde=None, base=None, metricas=None):
    """Actualizar y devolver una tabla Polars sin cargarla completa en memoria.

    Ejemplo: datos = actualizar_planes(); print(datos.head(10).collect())
    Volver a llamar esta función actualiza el mismo archivo. No guarda contraseñas.
    """
    if (not isinstance(dsn, str) or not dsn or dsn != dsn.strip()
            or any(char in dsn for char in ';{}\r\n\0')):
        raise ErrorDatos('El nombre ODBC debe ser el nombre existente, sin llaves ni parámetros adicionales.')
    sources = pyodbc.dataSources()
    dsn = next((name for name in sources if name.casefold() == dsn.casefold()), dsn)
    usuario = usuario or input('Usuario de la base de datos: ').strip()
    contrasena = contrasena or getpass('Contraseña (no se verá al escribir): ')
    if not usuario or not contrasena:
        raise ErrorDatos('Falta el usuario o la contraseña.')
    def escaped(value):
        return '{' + str(value).replace('}', '}}') + '}'
    connection_string = f'DSN={dsn};UID={escaped(usuario)};PWD={escaped(contrasena)}'
    _avance('[1/3] Conectando por ODBC en modo de solo lectura...')
    started = perf_counter()
    metrics = metricas if metricas is not None else {}
    with closing(pyodbc.connect(connection_string, readonly=True, autocommit=False, timeout=10)) as connection:
        metrics['conexion_s'] = perf_counter() - started
        _avance(f'[2/3] Conexión ODBC establecida ({metrics["conexion_s"]:.1f} s).')
        with closing(connection.cursor()) as cursor:
            target = extraer_y_guardar(cursor, salida, lote, ano_desde=ano_desde, base=base, metricas=metrics)
    return pl.scan_parquet(target)


def _ubicacion_tabla(path, table_name):
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(path) as archive:
        for part in archive.namelist():
            if part.startswith('xl/tables/') and part.endswith('.xml'):
                table = ET.fromstring(archive.read(part))
                if table.attrib.get('name') == table_name:
                    break
        else:
            raise ErrorDatos(f'El Excel no contiene la tabla {table_name}.')
        worksheet = None
        for name in archive.namelist():
            if name.startswith('xl/worksheets/_rels/') and name.endswith('.rels'):
                for rel in ET.fromstring(archive.read(name)):
                    target = rel.attrib['Target']
                    resolved = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl/worksheets', target))
                    if resolved == part:
                        worksheet = name.replace('/_rels/', '/').removesuffix('.rels')
        relationships = {}
        for rel in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels')):
            target = rel.attrib['Target']
            relationships[rel.attrib['Id']] = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl', target))
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        sheet = next((s.attrib['name'] for s in workbook.findall('s:sheets/s:sheet', ns)
            if relationships[s.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']] == worksheet), None)
        if sheet is None:
            raise ErrorDatos(f'No se encontró la hoja correspondiente a {table_name}.')
        columns = [c.attrib['name'] for c in table.findall('s:tableColumns/s:tableColumn', ns)]
        return sheet, table.attrib['ref'], columns, int(table.attrib.get('totalsRowCount', '0'))


def leer_excel(path):
    """Leer las tres tablas y los valores cacheados, sin recalcular ni guardar el Excel."""
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries
    from openpyxl.utils.escape import unescape

    path = Path(path).expanduser().resolve()
    if not path.is_file():
        raise ErrorDatos('No se encontró Despliegue 2022.xlsx. Indique su ruta en el parámetro excel.')
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=True)
    tables = {}
    try:
        def value(column, item):
            if isinstance(item, str):
                item = unescape(item)
            if column in {'FECINICIAL', 'FECFINAL'} and item is not None:
                if isinstance(item, (int, float)) and not isinstance(item, bool):
                    item = datetime(1899, 12, 30) + timedelta(days=int(item))
                elif isinstance(item, str):
                    try:
                        item = datetime.fromisoformat(item)
                    except ValueError:
                        item = datetime.strptime(item, '%d/%m/%Y')
                if not isinstance(item, (date, datetime)):
                    raise ErrorDatos(f'Fecha Excel inválida en {column}.')
                item = datetime.combine(item.date() if isinstance(item, datetime) else item, time.min)
            return item

        for name, excel_table in ((ESTRATEGIAS, 'Tabla1'), (MATRIZ, 'Matriz'), (CORPORATIVAS, 'Estrategias')):
            sheet, bounds, columns, totals = _ubicacion_tabla(path, excel_table)
            left, top, right, bottom = range_boundaries(bounds)
            rows = [tuple(value(c, v) for c, v in zip(columns, row))
                for row in workbook[sheet].iter_rows(min_row=top + 1, max_row=bottom - totals,
                    min_col=left, max_col=right, values_only=True)]
            dates = {c: pl.Datetime('ns') for c in columns if c in {'FECINICIAL', 'FECFINAL'}}
            frame = pl.DataFrame(rows, schema=columns, schema_overrides=dates,
                orient='row', strict=False, infer_schema_length=None)
            tables[name] = convertir_modelo(frame, name, desde_excel=True)
    finally:
        workbook.close()
    return tables


def _lazy(table):
    return table.lazy() if isinstance(table, pl.DataFrame) else table


def _sum(expression):
    return pl.when(expression.count() > 0).then(expression.sum()).otherwise(None)


def _count(expression):
    return pl.when(expression.count() > 0).then(expression.count().cast(pl.Int64)).otherwise(None)


def _divide(numerator, denominator, alternate=None):
    return pl.when(denominator.is_null() | (denominator == 0)).then(pl.lit(alternate, dtype=pl.Float64)).otherwise(numerator / denominator)


def _group(frame, columns, expressions):
    schema = frame.collect_schema()
    keys = [f'__grupo_{i}' for i in range(len(columns))]
    normalized = [
        (pl.col(c).str.to_lowercase() if schema[c] == pl.String else pl.col(c)).alias(k)
        for c, k in zip(columns, keys)]
    return frame.with_columns(normalized).group_by(keys, maintain_order=True).agg(
        [pl.col(c).first().alias(c) for c in columns] + expressions).drop(keys)


def _calendario(start, end):
    if start is None or end is None:
        return pl.DataFrame(schema={'Date': pl.Datetime('ns'), 'Año': pl.Int64, 'NroMes': pl.Int64,
            'Mes': pl.String, 'NroTrimestre': pl.Int64, 'Trimestre': pl.String, 'Día': pl.Int64}).lazy()
    days = pl.datetime_range(start, end, interval='1d', time_unit='ns', eager=True).alias('Date')
    months = dict(enumerate(('enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio',
        'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre'), 1))
    column = pl.col('Date')
    return pl.DataFrame(days).lazy().with_columns(
        column.dt.year().cast(pl.Int64).alias('Año'), column.dt.month().cast(pl.Int64).alias('NroMes'),
        column.dt.month().replace_strict(months, return_dtype=pl.String).alias('Mes'),
        column.dt.quarter().cast(pl.Int64).alias('NroTrimestre'),
        (pl.lit('Trim. ') + column.dt.quarter().cast(pl.String)).alias('Trimestre'),
        column.dt.day().cast(pl.Int64).alias('Día'))


def completar_modelo(tablas):
    """Generar las ocho tablas DAX una vez por actualización, sin filtros de consulta."""
    tables = {name: _lazy(table) for name, table in tablas.items()}
    applicable = pl.col('Filtro').str.to_lowercase() == 'aplica'
    active = applicable & (pl.col('ESTADO').str.to_lowercase() == 'a')
    summary = _group(tables[PLAN], ['PROCESO', 'Trimestre'], [
        pl.col('CALIFICACION').filter(applicable).mean().alias('CALIFICAIÓN'),
        _count(pl.col('TAREAS').filter(active)).alias('Tareas'),
        _count(pl.col('CALIFICACION').filter(active)).alias('TAREAS CALIFICADAS')])
    tables['Resumen'] = summary.with_columns(
        pl.when(pl.col('Tareas').eq_missing(pl.col('TAREAS CALIFICADAS')))
        .then(pl.lit('Completos')).otherwise(pl.lit('Incompletos')).alias('Validación'))
    tables['Resumen Estrategias'] = _group(tables[ESTRATEGIAS], ['ESTRATEG'], [
        _sum(pl.col('valor')).alias('cal'), _sum(pl.col('Peso')).alias('peso')]).with_columns(
        _divide(pl.col('cal'), pl.col('peso')).alias('Resultado'),
        (pl.col('cal') / pl.col('peso')).alias('prueba'))
    template = 'DateTableTemplate_4ef83330-4563-4c82-a681-915eb9ec9f36'
    tables[template] = _calendario(datetime(2015, 1, 1), datetime(2015, 1, 1))
    for name, (source, column) in CALENDARIOS.items():
        start, end = tables[source].select(pl.col(column).min().alias('inicio'),
            pl.col(column).max().alias('fin')).collect(engine='streaming').row(0)
        tables[name] = _calendario(datetime(start.year, 1, 1) if start else None,
                                  datetime(end.year, 12, 31) if end else None)
    return tables


class ModeloPowerBI:
    """Tablas Parquet diferidas y medidas con las relaciones originales.

    Filtros: {'Planes Operativos': {'ANO': [2026], 'Filtro': ['Aplica']}}.
    Una lista vacía selecciona cero filas. None selecciona BLANK.
    Las tablas Resumen son desconectadas, como en el PBIX.
    """
    def __init__(self, origen='modelo_powerbi'):
        self.carpeta = None
        if isinstance(origen, dict):
            self._tablas = {name: _lazy(table) for name, table in origen.items()}
        else:
            self.carpeta = Path(origen).expanduser().resolve()
            pointer = self.carpeta / 'actual.json'
            if not pointer.is_file():
                raise ErrorDatos('No hay una carga completa. Ejecute actualizar_modelo primero.')
            manifest = json.loads(pointer.read_text(encoding='utf-8'))
            self._tablas = {}
            for name, relative in manifest['tablas'].items():
                path = (self.carpeta / relative).resolve()
                if not path.is_relative_to(self.carpeta) or not path.is_file():
                    raise ErrorDatos('La carga contiene una ruta inválida o un archivo ausente.')
                self._tablas[name] = pl.scan_parquet(path)
        self.tablas = tuple(self._tablas)

    def tabla(self, nombre, filtros=None):
        if nombre not in self._tablas:
            raise ErrorDatos(f'Tabla no inventariada: {nombre}')
        context = filtros or {}
        for name, selections in context.items():
            if name not in self._tablas or not isinstance(selections, dict):
                raise ErrorDatos('Los filtros deben indicar una tabla y sus columnas.')
            if set(selections) - set(self._tablas[name].collect_schema()):
                raise ErrorDatos(f'El filtro contiene columnas desconocidas de {name}.')
        filtered, affected = {}, {}

        def has_filter(name):
            if name not in affected:
                affected[name] = bool(context.get(name)) or any(has_filter(parent)
                    for parent, _, child, _ in RELACIONES if child == name)
            return affected[name]

        def visit(name):
            if name in filtered:
                return filtered[name]
            table = self._tablas[name]
            schema = table.collect_schema()
            for column, selection in context.get(name, {}).items():
                values = list(selection) if isinstance(selection, (list, tuple, set)) else [selection]
                expression = pl.col(column)
                if schema[column] == pl.String:
                    expression = expression.str.to_lowercase()
                    values = [v.lower() if isinstance(v, str) else v for v in values]
                nonnull = [v for v in values if v is not None]
                condition = expression.is_in(nonnull).fill_null(False) if nonnull else pl.lit(False)
                if None in values:
                    condition = condition | expression.is_null()
                table = table.filter(condition)
            for parent, parent_key, child, child_key in RELACIONES:
                if child != name or not has_filter(parent):
                    continue
                dimension = visit(parent)
                left, right = pl.col(child_key), pl.col(parent_key)
                if schema[child_key] == pl.String:
                    left, right = left.str.to_lowercase(), right.str.to_lowercase()
                dimension = dimension.select(right.alias('__relacion')).unique()
                table = table.with_columns(left.alias('__relacion')).join(dimension,
                    on='__relacion', how='semi', nulls_equal=True).drop('__relacion')
            filtered[name] = table
            return table
        return visit(nombre)

    def _planes(self, filtros):
        row = self.tabla(PLAN, filtros).select(pl.col('TAREAS').count().alias('tareas'),
            pl.col('CALIFICACION').mean().alias('nota'), pl.col('Trimestre').n_unique().alias('n'),
            pl.col('Trimestre').first().alias('trimestre')).collect(engine='streaming').row(0, named=True)
        quarter = row['trimestre']
        meta = {k.lower(): v for k, v in METAS.items()}.get(quarter.lower()) if row['n'] == 1 and quarter else None
        compliance = 0. if meta is None else row['nota'] / meta if row['nota'] is not None else None
        # Literal autorizado: el filtro Estrategias[ANO]=2024 no retropropaga a Planes.
        return {'Meta Trim': meta, '% Cumplimiento': compliance, 'Tareas 2024': row['tareas'] - 2}

    def medida(self, nombre, filtros=None):
        if nombre == 'prueba':
            raise ErrorDatos('Planes Operativos[prueba] está sin fórmula en el PBIX; permanece bloqueada.')
        if nombre in {'Meta Trim', '% Cumplimiento', 'Tareas 2024'}:
            return self._planes(filtros)[nombre]
        if nombre == 'Hora Actualización':
            value = self.tabla('Fecha Actualización', filtros).select(pl.col('Fecha Actualización').max()).collect().item()
            return value - timedelta(hours=5) if value is not None else None
        if nombre in {'Resultado', 'Resultado2'}:
            table, num, den = (ESTRATEGIAS, 'valor', 'Peso') if nombre == 'Resultado' else (MATRIZ, 'VALOR', 'PESO2')
            return self.tabla(table, filtros).select(_divide(_sum(pl.col(num)), _sum(pl.col(den)))).collect(engine='streaming').item()
        raise ErrorDatos(f'Medida no inventariada: {nombre}')

    def medidas(self, filtros=None):
        return {**self._planes(filtros), **{name: self.medida(name, filtros)
            for name in ('Resultado', 'Resultado2', 'Hora Actualización')}}

    def agrupar(self, tabla, columnas, filtros=None):
        frame = self.tabla(tabla, filtros)
        columnas = [columnas] if isinstance(columnas, str) else list(columnas)
        if not columnas or set(columnas) - set(frame.collect_schema()):
            raise ErrorDatos('Indique columnas existentes para agrupar.')
        if tabla == PLAN:
            meta = pl.when(pl.col('Trimestre').n_unique() == 1).then(
                pl.col('Trimestre').first().str.to_lowercase().replace_strict(
                    {k.lower(): v for k, v in METAS.items()}, default=None, return_dtype=pl.Float64)).otherwise(None)
            return _group(frame, columnas, [pl.col('CALIFICACION').mean().alias('Calificación'),
                _count(pl.col('TAREAS')).alias('Tareas'), _count(pl.col('CALIFICACION')).alias('Calificadas'),
                pl.col('CALIFICACION').is_null().sum().alias('Pendientes'),
                pl.col('ACTIVID').str.to_lowercase().n_unique().alias('Actividades'), meta.alias('Meta Trim'),
                (pl.col('TAREAS').count().cast(pl.Int64) - 2).alias('Tareas 2024')]).with_columns(
                    _divide(pl.col('Calificación'), pl.col('Meta Trim'), 0.).alias('% Cumplimiento'))
        if tabla in {ESTRATEGIAS, MATRIZ}:
            num, den, result = ('valor', 'Peso', 'Resultado') if tabla == ESTRATEGIAS else ('VALOR', 'PESO2', 'Resultado2')
            return _group(frame, columnas, [pl.len().alias('Registros'),
                _divide(_sum(pl.col(num)), _sum(pl.col(den))).alias(result)])
        raise ErrorDatos('Los agregados corresponden a Planes Operativos, Estrategias o Matriz Objetivos.')


def _excel_original(excel):
    if excel is not None:
        path = Path(excel).expanduser().resolve()
    else:
        script = Path(globals().get('__file__', Path.cwd() / 'actualizar_planes.py')).resolve().parent
        candidates = [root / name for root in dict.fromkeys((Path.cwd(), script))
            for name in ('Despliegue 2022.xlsx', 'Tema tablero institucional/Despliegue 2022.xlsx')]
        path = next((p for p in candidates if p.is_file()), None)
        if path is None:
            path = Path(input('Ruta de Despliegue 2022.xlsx: ').strip().strip('"')).expanduser().resolve()
    if not path.is_file():
        raise ErrorDatos('No se encontró el Excel original requerido por el Power BI.')
    return path


def _ventana_abierta(previous, now):
    current = now.astimezone(COLOMBIA).year
    year = previous.get('ultimo_anio_abierto')
    if year is None:
        try:
            stamp = datetime.fromisoformat(previous['actualizado_utc'])
            # Las cargas originales guardaban UTC sin sufijo de zona.
            year = stamp.replace(tzinfo=stamp.tzinfo or timezone.utc).astimezone(COLOMBIA).year
        except (KeyError, TypeError, ValueError):
            raise ErrorDatos('El corte anterior no tiene fecha válida. Usa --reconstruir para reconstruir el histórico.') from None
    if type(year) is not int or not 1 <= year <= current:
        raise ErrorDatos('El año del corte anterior es inválido o futuro. Revisa el reloj y el corte.')
    return min(current, year)


def _firma_transformacion():
    text = '\n'.join(inspect.getsource(fn) for fn in (transform_operational, convertir_modelo))
    text += repr((ESQUEMA, AREA_REPLACEMENTS, PROCESS_REPLACEMENTS, FINAL_PROCESS))
    return hashlib.sha256(text.encode('utf-8')).hexdigest()


def _validar_base(root, previous):
    expected = {*ESQUEMAS, 'Metas', 'Trimestres', 'Fecha Actualización', 'Resumen',
                'Resumen Estrategias', 'DateTableTemplate_4ef83330-4563-4c82-a681-915eb9ec9f36', *CALENDARIOS}
    try:
        if set(previous['tablas']) != expected:
            raise ValueError('tablas incompletas')
        for name, relative in previous['tablas'].items():
            path = (root / relative).resolve()
            if not path.is_relative_to(root) or not path.is_file():
                raise ValueError('ruta inválida')
            frame = pl.scan_parquet(path)
            if name in ESQUEMAS and frame.collect_schema() != pl.Schema(ESQUEMAS[name]):
                raise ValueError('esquema incompatible')
            frame.select(pl.all().null_count()).collect()  # Verificar lectura de todos los bloques.
        if previous.get('firma_transformacion', _firma_transformacion()) != _firma_transformacion():
            raise ValueError('transformación incompatible')
    except Exception:
        raise ErrorDatos('La base guardada está incompleta o es incompatible. Se conserva; usa --reconstruir para reconstruirla.') from None
    return (root / previous['tablas'][PLAN]).resolve()


@contextmanager
def _bloqueo_actualizacion(root):
    """Bloqueo del sistema operativo: se libera incluso si muere el proceso."""
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    with (root / '.actualizacion.lock').open('a+b') as handle:
        if handle.tell() == 0:
            handle.write(b'0')
            handle.flush()
        handle.seek(0)
        try:
            if os.name == 'nt':
                import msvcrt
                msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError:
            raise ErrorDatos('Ya hay una actualización en curso en esta carpeta.') from None
        try:
            yield
        finally:
            handle.seek(0)
            if os.name == 'nt':
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def actualizar_modelo(excel=None, dsn='salud', usuario=None, contrasena=None,
                      salida='modelo_powerbi', lote=10000, *, reconstruir=False):
    with _bloqueo_actualizacion(Path(salida).expanduser().resolve()):
        return _actualizar_modelo(excel, dsn, usuario, contrasena, salida, lote, reconstruir=reconstruir)


def _actualizar_modelo(excel=None, dsn='salud', usuario=None, contrasena=None,
                       salida='modelo_powerbi', lote=10000, *, reconstruir=False):
    """Actualizar las 15 tablas y publicar el conjunto completo en una sola operación.

    Conserva la versión anterior y no publica una actualización incompleta.
    Para trabajar sin volver a consultar: ModeloPowerBI(salida).
    """
    started = perf_counter()
    now_utc = datetime.now(timezone.utc)
    root = Path(salida).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    pointer = root / 'actual.json'
    previous_bytes = pointer.read_bytes() if pointer.is_file() else None
    try:
        previous = json.loads(previous_bytes) if previous_bytes else {}
    except (ValueError, UnicodeError):
        if not reconstruir:
            raise ErrorDatos('El corte guardado es ilegible. Usa --reconstruir; se conservan los datos anteriores.') from None
        previous = {}
    base = ano_desde = None
    if previous_bytes is not None and not reconstruir:
        base = _validar_base(root, previous)
        ano_desde = _ventana_abierta(previous, now_utc)
    metrics = {'modo': 'incremental' if base else 'completa', 'ano_desde': ano_desde}
    excel = _excel_original(excel)
    _avance('Leyendo las tres tablas del Excel...')
    imported = leer_excel(excel)
    _avance('Lectura del Excel terminada.')
    version = uuid.uuid4().hex
    versions = root / 'versiones'
    versions.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.cargando-', dir=root) as directory:
        work = Path(directory).resolve()
        if not work.is_relative_to(root):
            raise ErrorDatos('Directorio temporal inválido.')
        operational = work / 'operativo.parquet'
        # La espera de credenciales no forma parte de las mediciones.
        local_before_credentials = perf_counter() - started
        usuario = usuario or input('Usuario de la base de datos: ').strip()
        contrasena = contrasena or getpass('Contraseña (no se verá al escribir): ')
        ready = perf_counter()
        actualizar_planes(dsn, usuario, contrasena, operational, lote,
                          ano_desde=ano_desde, base=base, metricas=metrics)
        goals = pl.DataFrame({'Trimestre': TRIMESTRES, 'Meta': list(METAS.values())})
        now = now_utc.replace(tzinfo=None)
        _avance('Calculando resúmenes y calendarios del modelo...')
        tables = completar_modelo({PLAN: pl.scan_parquet(operational), **imported,
            'Metas': goals, 'Trimestres': goals.clone(),
            'Fecha Actualización': pl.DataFrame({'Fecha Actualización': [now]},
                schema={'Fecha Actualización': pl.Datetime('ns')})})
        # La unicidad del lado uno es necesaria para conservar las relaciones M:1.
        for parent, key, _, _ in RELACIONES:
            column = pl.col(key)
            if tables[parent].collect_schema()[key] == pl.String:
                column = column.str.to_lowercase()
            rows, distinct = tables[parent].select(pl.len(), column.n_unique().alias('unicos')).collect().row(0)
            if rows != distinct:
                raise ErrorDatos(f'La relación requiere valores únicos en {parent}[{key}]. Se conserva la carga anterior.')
        manifest = {'version': version, 'anterior': previous.get('version'),
            'actualizado_utc': now_utc.isoformat(), 'excel': excel.name, 'tablas': {},
            'ultimo_anio_abierto': now_utc.astimezone(COLOMBIA).year,
            'firma_transformacion': _firma_transformacion(), 'actualizacion': metrics,
            'medidas_bloqueadas': ['Planes Operativos[prueba]']}
        _avance('Guardando el modelo completo...')
        for index, (name, table) in enumerate(tables.items()):
            if name == PLAN:
                continue
            filename = f'tabla_{index:02d}.parquet'
            table.sink_parquet(work / filename, compression='zstd', engine='streaming')
            manifest['tablas'][name] = f'versiones/{version}/{filename}'
        # Mover el operativo ya escrito evita copiar y comprimir de nuevo el hecho mayor.
        os.replace(operational, work / 'tabla_00.parquet')
        manifest['tablas'][PLAN] = f'versiones/{version}/tabla_00.parquet'
        destination = versions / version
        os.replace(work, destination)
        temporary_pointer = root / f'.actual-{version}.json'
        try:
            # Leer las 15 tablas candidatas antes de cambiar el corte visible.
            _validar_base(root, manifest)
            if (pointer.read_bytes() if pointer.is_file() else None) != previous_bytes:
                raise ErrorDatos('El corte cambió durante la actualización. Se conserva el corte publicado.')
            metrics['total_s'] = perf_counter() - ready + local_before_credentials
            metrics['procesamiento_local_s'] = metrics['total_s'] - sum(metrics[k] for k in ('conexion_s', 'consulta_s', 'descarga_s'))
            temporary_pointer.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
            os.replace(temporary_pointer, pointer)
        except BaseException:
            if destination.is_relative_to(versions) and destination.name == version:
                shutil.rmtree(destination)
            temporary_pointer.unlink(missing_ok=True)
            raise
    # Conservar las versiones existentes; también pueden estar siendo leídas por el tablero.
    model = ModeloPowerBI(root)
    _avance(f'Modelo completo actualizado: {len(model.tablas)} tablas. Carpeta: {root}')
    _avance('Tiempos (s): ' + '; '.join(f'{key}={metrics[key]:.3f}' for key in
            ('conexion_s', 'consulta_s', 'descarga_s', 'procesamiento_local_s', 'total_s')))
    return model


def credenciales_locales():
    """Ingreso local opcional, sin argumentos de proceso, archivos ni chat."""
    import tkinter as tk
    from tkinter import simpledialog
    window = tk.Tk()
    window.withdraw()
    window.attributes('-topmost', True)
    try:
        user = simpledialog.askstring('Actualizar planes — IBM i',
            'Usuario del DSN salud (solo se utiliza en memoria):', parent=window)
        if not user:
            raise ErrorDatos('Actualización cancelada sin consultar IBM i.')
        password = simpledialog.askstring('Actualizar planes — IBM i',
            'Contraseña (no se guarda ni se muestra):', show='*', parent=window)
        if not password:
            raise ErrorDatos('Actualización cancelada sin consultar IBM i.')
        return user.strip(), password
    finally:
        window.destroy()


def main():
    parser = argparse.ArgumentParser(description='Actualizar el modelo local mediante SELECT de solo lectura.')
    parser.add_argument('--reconstruir', action='store_true', help='Carga completa excepcional, también del histórico cerrado.')
    parser.add_argument('--credenciales-ventana', action='store_true', help='Solicitar credenciales en ventanas locales, sin guardarlas.')
    args = parser.parse_args()
    try:
        global modelo
        user, password = credenciales_locales() if args.credenciales_ventana else (None, None)
        modelo = actualizar_modelo(usuario=user, contrasena=password, reconstruir=args.reconstruir)
        password = None
        for name, value in modelo.medidas().items():
            print(f'{name}: {value}')
        return 0
    except pyodbc.Error as error:
        hints = {
            'IM002': 'Windows no encuentra el DSN ODBC de 64 bits.',
            'IM003': 'Windows no pudo cargar el controlador ODBC.',
            'IM014': 'Python y el controlador ODBC tienen distinta arquitectura.',
            '28000': 'El origen rechazó el usuario o la contraseña.',
            '08001': 'No se pudo conectar al servidor definido en el DSN.',
            '08S01': 'Se interrumpió la comunicación con IBM i.',
            'HYT00': 'Se agotó el tiempo de conexión.',
            'HYT01': 'Se agotó el tiempo de conexión.',
            '42000': 'El origen rechazó la consulta; revisar acceso a la vista.',
        }
        state = next((arg for arg in error.args if isinstance(arg, str) and arg in hints), None)
        print(f'Error ODBC {state}: {hints[state]}' if state else 'El controlador ODBC no pudo completar la lectura.')
    except ErrorDatos as error:
        print(str(error))
    except KeyboardInterrupt:
        print('Actualización cancelada.')
    except Exception as error:
        print(f'No terminó la actualización ({type(error).__name__}). Revisar datos y carpeta de salida.')
    return 1


if __name__ == '__main__':
    sys.exit(main())
